"""Tests de caractérisation du DocumentLoader : un fichier minimal par format, plus les métadonnées.

Les fichiers sont générés dans `tmp_path` (fixture `data_dir`). Les comportements
connus pour être discutables sont figés par un test normal ; celui qui décrit le
comportement voulu est un `xfail(strict=True)` rattaché au lot qui le corrige.
"""

import email.message
import logging
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from unittest.mock import patch

import pytest
from docx import Document as DocxDocument
from PIL import Image

from conftest import write_files
from document_loader import DocumentLoader


def make_pdf(path, text: str) -> None:
    """PDF minimal d'une page, écrit à la main (pas de dépendance supplémentaire)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = "%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.write_bytes(out.encode("latin-1"))


@pytest.fixture
def loader(data_dir):
    return DocumentLoader()


def load_one(loader, name: str):
    docs = loader.load_specific([name])
    assert len(docs) == 1, f"{name} n'a pas été chargé"
    return docs[0]


# --- Un fichier minimal par format --------------------------------------------

class TestFormats:
    def test_txt_utf8(self, loader, data_dir):
        text = "Crème brûlée à la vanille, préparation du dessert familial"
        write_files(data_dir, {"note.txt": text})
        doc = load_one(loader, "note.txt")
        assert doc.page_content == text
        assert doc.metadata["file_type"] == "txt"

    def test_txt_latin1_encoding_is_detected(self, loader, data_dir):
        (data_dir / "vieux.txt").write_bytes("Crème brûlée à la vanille, préparation du dessert familial".encode("latin-1"))
        doc = load_one(loader, "vieux.txt")
        assert "Crème brûlée" in doc.page_content

    def test_markdown_is_loaded_as_text(self, loader, data_dir):
        write_files(data_dir, {"liste.md": "# Courses\n\n- pommes\n- poires\n"})
        doc = load_one(loader, "liste.md")
        assert doc.page_content == "# Courses\n\n- pommes\n- poires\n"
        assert doc.metadata["file_type"] == "md"

    def test_pdf_text_is_extracted(self, loader, data_dir):
        make_pdf(data_dir / "releve.pdf", "Releve bancaire de mars")
        doc = load_one(loader, "releve.pdf")
        assert "Releve bancaire de mars" in doc.page_content
        assert doc.metadata["file_type"] == "pdf"

    def test_docx_paragraphs_are_joined_by_blank_lines(self, loader, data_dir):
        document = DocxDocument()
        document.add_paragraph("Premier paragraphe")
        document.add_paragraph("")
        document.add_paragraph("Second paragraphe")
        document.save(data_dir / "lettre.docx")
        doc = load_one(loader, "lettre.docx")
        assert doc.page_content == "Premier paragraphe\n\nSecond paragraphe"

    def test_docx_tables_are_ignored(self, loader, data_dir):
        """Caractérisation de F16 : le contenu des tableaux DOCX est perdu (lot 11)."""
        document = DocxDocument()
        document.add_paragraph("Avant le tableau")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "loyer"
        table.cell(0, 1).text = "800"
        document.save(data_dir / "budget.docx")
        doc = load_one(loader, "budget.docx")
        assert doc.page_content == "Avant le tableau"

    def test_eml_headers_and_plain_body(self, loader, data_dir):
        message = email.message.EmailMessage()
        message["Subject"] = "Diner dimanche"
        message["From"] = "Marie Dupont <marie@example.org>"
        message["To"] = "papa@example.org"
        message["Date"] = "Sat, 02 Mar 2024 10:00:00 +0000"
        message.set_content("On mange chez nous dimanche.")
        (data_dir / "invitation.eml").write_bytes(bytes(message))
        doc = load_one(loader, "invitation.eml")
        assert doc.page_content.startswith("Sujet: Diner dimanche\nDe: Marie Dupont <marie@example.org>\n")
        assert "On mange chez nous dimanche." in doc.page_content

    def test_eml_multipart_with_html_only_has_an_empty_body(self, loader, data_dir):
        """Caractérisation de F15 : seule la partie text/plain est lue (lot 11)."""
        message = MIMEMultipart("alternative")
        message["Subject"] = "Newsletter"
        message.attach(MIMEText("<p>Contenu HTML</p>", "html"))
        (data_dir / "news.eml").write_bytes(message.as_bytes())
        doc = load_one(loader, "news.eml")
        assert "Contenu HTML" not in doc.page_content
        assert "Sujet: Newsletter" in doc.page_content

    @pytest.mark.xfail(reason="F23 : l'en-tête Subject encodé (RFC 2047) doit être décodé (lot 11)")
    def test_eml_encoded_subject_is_decoded(self, loader, data_dir):
        raw = (
            b"Subject: =?utf-8?q?D=C3=AEner_dimanche?=\r\n"
            b"From: marie@example.org\r\n"
            b"Content-Type: text/plain; charset=utf-8\r\n\r\n"
            b"On mange chez nous.\r\n"
        )
        (data_dir / "encode.eml").write_bytes(raw)
        doc = load_one(loader, "encode.eml")
        assert doc.page_content.startswith("Sujet: Dîner dimanche\n")

    @pytest.mark.xfail(reason="F24 : un court fichier UTF-8 ne doit pas être lu en ISO-8859-1 (lot 11)")
    def test_txt_short_utf8_is_not_misdetected_as_latin1(self, loader, data_dir):
        text = "Bonjour à tous, voici une note."
        (data_dir / "court.txt").write_bytes(text.encode("utf-8"))
        assert load_one(loader, "court.txt").page_content == text

    def test_image_goes_through_ocr_with_french_and_english(self, loader, data_dir):
        Image.new("RGB", (20, 20), "white").save(data_dir / "scan.png")
        with patch("document_loader.pytesseract.image_to_string", return_value="  Texte reconnu \n") as ocr:
            doc = load_one(loader, "scan.png")
        assert doc.page_content == "Texte reconnu"
        assert ocr.call_args.kwargs["lang"] == "fra+eng"
        assert doc.metadata["file_type"] == "png"


# --- Découverte et erreurs ----------------------------------------------------

class TestDiscovery:
    def test_list_files_returns_relative_paths_for_supported_extensions_sorted_by_name(self, loader, data_dir):
        write_files(data_dir, {"b/zeta.txt": "z", "a/Alpha.md": "a", "ignore.xyz": "x", "gamma.txt": "g"})
        files = loader.list_files()
        assert [f["name"] for f in files] == ["Alpha.md", "gamma.txt", "zeta.txt"]
        assert [f["path"] for f in files] == ["a/Alpha.md", "gamma.txt", "b/zeta.txt"]
        assert files[0]["type"] == "md"

    def test_load_all_loads_every_supported_file(self, loader, data_dir):
        write_files(data_dir, {"a.txt": "alpha", "sous/b.md": "beta", "c.xyz": "ignoré"})
        docs = loader.load_all()
        assert sorted(d.page_content for d in docs) == ["alpha", "beta"]

    def test_empty_file_is_skipped(self, loader, data_dir):
        write_files(data_dir, {"vide.txt": ""})
        assert loader.load_all() == []
        assert loader.load_specific(["vide.txt"]) == []

    def test_missing_unsupported_and_directory_paths_are_skipped(self, loader, data_dir):
        write_files(data_dir, {"x.xyz": "inconnu", "dossier/a.txt": "contenu"})
        assert loader.load_specific(["absent.txt", "x.xyz", "dossier"]) == []

    def test_corrupt_file_is_logged_by_load_specific_only(self, loader, data_dir, caplog):
        """Caractérisation de F18 : load_all ignore les erreurs sans rien journaliser."""
        (data_dir / "casse.pdf").write_bytes(b"ceci n'est pas un pdf")
        with caplog.at_level(logging.ERROR, logger="document_loader"):
            assert loader.load_all() == []
            assert caplog.records == []
            assert loader.load_specific(["casse.pdf"]) == []
        assert any("casse.pdf" in record.getMessage() for record in caplog.records)


# --- Métadonnées --------------------------------------------------------------

class TestMetadata:
    def test_metadata_keys_for_a_complete_document(self, loader, data_dir):
        text = "Facture n° 12\nPar: Marie Dupont\nMontant TTC : 120 euros, TVA incluse."
        write_files(data_dir, {"facture_2024-03-15.txt": text})
        meta = load_one(loader, "facture_2024-03-15.txt").metadata
        assert meta == {
            "source": "facture_2024-03-15.txt",
            "file_type": "txt",
            "size_bytes": len(text.encode("utf-8")),
            "date": "2024-03-15",
            "year": 2024,
            "doc_type": "facture",
            "author": "Marie Dupont",
        }

    def test_author_is_absent_when_not_found(self, loader, data_dir):
        write_files(data_dir, {"memo.txt": "Rappel : sortir les poubelles."})
        assert "author" not in load_one(loader, "memo.txt").metadata

    def test_date_and_year_fall_back_on_content_then_file_date(self, loader, data_dir):
        write_files(data_dir, {"a.txt": "Signé le 12 mars 2021 à Lyon", "b.txt": "Rien de daté ici"})
        meta_a = load_one(loader, "a.txt").metadata
        assert meta_a["date"] == "12 mars 2021"
        assert meta_a["year"] == 2021
        meta_b = load_one(loader, "b.txt").metadata
        assert meta_b["date"].startswith(str(meta_b["year"]))  # date de modification du fichier

    def test_classification_by_keywords(self, loader):
        assert loader._classify_document("Ingrédients : farine. Cuisson au four.") == "recette"
        assert loader._classify_document("Objet: rendez-vous\nDe: Paul") == "email"
        assert loader._classify_document("zzz") == "autre"

    def test_classification_matches_substrings(self, loader):
        """Caractérisation de F17 : « ml » est trouvé dans « html » (lot 11)."""
        assert loader._classify_document("<html></html>") == "recette"

    def test_source_is_the_bare_file_name(self, loader, data_dir):
        """Caractérisation de F7 : deux fichiers homonymes ont la même source."""
        write_files(data_dir, {"a/facture.txt": "facture A", "b/facture.txt": "facture B"})
        docs = loader.load_specific(["a/facture.txt", "b/facture.txt"])
        assert [d.metadata["source"] for d in docs] == ["facture.txt", "facture.txt"]

    @pytest.mark.xfail(reason="F7/D5 : source = chemin relatif à /data (lot 9)")
    def test_source_is_the_relative_path(self, loader, data_dir):
        write_files(data_dir, {"a/facture.txt": "facture A", "b/facture.txt": "facture B"})
        docs = loader.load_specific(["a/facture.txt", "b/facture.txt"])
        assert [d.metadata["source"] for d in docs] == ["a/facture.txt", "b/facture.txt"]
