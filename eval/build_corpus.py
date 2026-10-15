"""Génère eval/corpus/ : documents fictifs de la famille Martin (aucune donnée réelle).

Les fichiers produits sont versionnés ; ce script ne sert qu'à les régénérer :
    .venv/bin/python eval/build_corpus.py
"""

import sys
from email.message import EmailMessage
from pathlib import Path

from docx import Document as DocxDocument
from PIL import Image, ImageDraw, ImageFont

CORPUS = Path(__file__).parent / "corpus"


def write_pdf(path: Path, lines: list[str]) -> None:
    """PDF texte d'une page, écrit à la main (police Helvetica, encodage WinAnsi)."""
    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = "BT /F1 12 Tf 72 720 Td 16 TL\n" + "\n".join(f"({escape(line)}) '" for line in lines) + "\nET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = "%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{offset:010d} 00000 n \n" for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(out.encode("latin-1"))


def build() -> None:
    CORPUS.mkdir(exist_ok=True)

    write_pdf(CORPUS / "assurance" / "contrat_habitation.pdf", [
        "Contrat d'assurance habitation n° HAB-2023-0417",
        "Assuré : famille Martin, 12 rue des Lilas, 69003 Lyon",
        "Assureur : Mutuelle du Rhône",
        "Prime annuelle : 412 euros, payable en deux fois.",
        "Franchise : 250 euros par sinistre (dégât des eaux, vol, incendie).",
        "Date d'effet : 1er janvier 2023.",
        "Échéance du contrat : 31 décembre 2025, reconduction tacite.",
    ])
    write_pdf(CORPUS / "edf" / "facture.pdf", [
        "Facture d'électricité EDF - novembre 2024",
        "Client : famille Martin",
        "Consommation : 412 kWh sur la période.",
        "Montant total à payer : 87,40 euros.",
        "Date limite de paiement : 15 décembre 2024.",
    ])
    write_pdf(CORPUS / "garage" / "facture.pdf", [
        "Garage Dupuis - Facture de réparation",
        "Client : Marie Martin, véhicule Renault Clio",
        "Révision et remplacement des plaquettes de frein.",
        "Montant total à payer : 342,00 euros.",
        "Date de la facture : 8 mars 2024.",
    ])

    document = DocxDocument()
    document.add_paragraph("Budget familial 2024 - prévisions validées en janvier.")
    table = document.add_table(rows=4, cols=2)
    for row, (poste, montant) in enumerate([
        ("Poste", "Budget annuel (euros)"),
        ("Vacances", "1800"),
        ("Loisirs des enfants", "950"),
        ("Cadeaux de Noël", "600"),
    ]):
        table.cell(row, 0).text = poste
        table.cell(row, 1).text = montant
    document.add_paragraph("Le budget sera révisé en juillet.")
    document.save(CORPUS / "budget_2024.docx")

    invitation = EmailMessage()
    invitation["Subject"] = "Invitation anniversaire"
    invitation["From"] = "Marie Martin <marie@example.org>"
    invitation["To"] = "famille@example.org"
    invitation["Date"] = "Mon, 04 Mar 2024 09:00:00 +0000"
    invitation.set_content(
        "Bonjour à tous,\n\nL'anniversaire de Léa aura lieu samedi 16 mars à 15 h 30 à la maison.\n"
        "Tante Claire apporte le gâteau, Hugo s'occupe de la musique.\n"
        "Merci de confirmer votre présence avant jeudi.\n\nMarie\n",
        charset="utf-8",
    )
    (CORPUS / "invitation_anniversaire.eml").write_bytes(bytes(invitation))

    # Sujet encodé (RFC 2047) : écrit à la main pour contrôler l'en-tête brut (F23)
    (CORPUS / "sujet_encode.eml").write_bytes(
        b"Subject: =?utf-8?q?D=C3=AEner_dimanche?=\r\n"
        b"From: Marie Martin <marie@example.org>\r\n"
        b"To: famille@example.org\r\n"
        b"Date: Wed, 06 Mar 2024 18:00:00 +0000\r\n"
        b"Content-Type: text/plain; charset=utf-8\r\n\r\n"
        b"On se retrouve chez mamie pour le d\xc3\xaener de dimanche, vers 19 h.\r\n"
    )

    image = Image.new("RGB", (900, 260), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=34)
    draw.text((30, 30), "Tarte aux pommes de mamie", fill="black", font=font)
    draw.text((30, 100), "Four : 180 degres pendant 35 minutes", fill="black", font=font)
    draw.text((30, 170), "Laisser refroidir avant de servir", fill="black", font=font)
    image.save(CORPUS / "recette_scan.png")

    (CORPUS / "recettes.md").write_text(
        "# Recettes de famille\n\n## Tarte aux pommes\n\nIngrédients :\n\n"
        "- 6 pommes\n- 250 g de farine\n- 125 g de beurre\n- 80 g de sucre\n\n"
        "Étapes :\n\n1. Préparer la pâte et la laisser reposer.\n2. Disposer les pommes en rosace.\n"
        "3. Cuire au four.\n\n## Crêpes\n\n- 250 g de farine\n- 3 oeufs\n- 50 cl de lait\n",
        encoding="utf-8",
    )
    (CORPUS / "menage.md").write_text(
        "# Planning des tâches ménagères\n\n| Jour | Tâche | Responsable |\n|---|---|---|\n"
        "| Lundi | Vaisselle | Léa |\n| Mardi | Sortir les poubelles | Hugo |\n"
        "| Mercredi | Aspirateur | Marie |\n| Jeudi | Courses | Papa |\n",
        encoding="utf-8",
    )
    # Un seul caractère accentué : chardet le classe en ISO-8859-1 (confiance 0,73), voir F24
    (CORPUS / "note_courte.txt").write_text(
        "Bonjour à tous, la clef du chalet est dans le tiroir du couloir.\n", encoding="utf-8"
    )


if __name__ == "__main__":
    build()
    print(f"Corpus écrit dans {CORPUS}", file=sys.stderr)
